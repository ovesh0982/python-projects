from docx import Document

# Recreate the document since the execution environment was reset
doc = Document()

# Title
doc.add_heading('Ovesh Abdulaziz Basan', 0)
doc.add_paragraph('608, Al-Mujdalifa Apartment, Pathanwadi, Malad East, Mumbai - 400097')
doc.add_paragraph('Phone: 8879806032 | Email: oveshaziz123@gmail.com\nGitHub: https://github.com/ovesh0982')

doc.add_paragraph('')

# Career Objective
doc.add_heading('Career Objective', level=1)
doc.add_paragraph(
    "A passionate and detail-oriented BCA graduate with strong foundational knowledge in Python, "
    "data analysis, and web development. Seeking an entry-level position in the IT industry where I "
    "can utilize my skills in software development, data science, and automation to contribute effectively "
    "and grow professionally."
)

# Education
doc.add_heading('Educational Qualification', level=1)
doc.add_paragraph(
    "Bachelor of Computer Applications (BCA)\n"
    "Tilak Maharashtra Vidyapeeth, Pune\n"
    "Year of Passing: 2025 | CGPA: 8.0"
)
doc.add_paragraph(
    "HSC (12th) – [Board Name]\n"
    "Year of Passing: 2022 | Percentage: 55%"
)
doc.add_paragraph(
    "SSC (10th) – [Board Name]\n"
    "Year of Passing: 2020 | Percentage: 84%"
)

# Technical Skills
doc.add_heading('Technical Skills', level=1)
doc.add_paragraph(
    "- Programming Languages: C, Python, PHP\n"
    "- Web Technologies: HTML, CSS, JavaScript\n"
    "- Database: MySQL\n"
    "- Data Skills: Data Analysis, Data Science, Machine Learning Algorithms\n"
    "- Tools: Automation scripting, Git, VS Code"
)

# Projects
doc.add_heading('Projects', level=1)
doc.add_paragraph(
    "1. AI Voice Assistant (JARVIS)\n"
    "Developed a Python-based voice assistant with speech recognition and automation features for tasks like opening apps, searching the web, and more."
)
doc.add_paragraph(
    "2. Data Analysis Dashboard\n"
    "Created a dashboard using Pandas and Matplotlib to analyze and visualize large datasets for sales insights."
)
doc.add_paragraph(
    "3. Web-based Student Portal\n"
    "Built a student management portal using PHP, MySQL, HTML/CSS for data entry and academic records management."
)

# Languages
doc.add_heading('Languages Known', level=1)
doc.add_paragraph("English, Hindi")

# Hobbies
doc.add_heading('Hobbies', level=1)
doc.add_paragraph("Playing Cricket")

# Save the updated document
updated_file_path = "/mnt/data/Ovesh_Resume_With_GitHub.docx"
doc.save(updated_file_path)

updated_file_path
